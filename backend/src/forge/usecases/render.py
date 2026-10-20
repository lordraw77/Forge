"""Output files of a run (M2.3d): DOCX, XLSX and PDF generated from the final output, and the
text extracted from the files an end user attaches to a form.

Pure functions over bytes, no I/O. The final output is Markdown text or a JSON value; a small
block parser (headings, paragraphs, bullet and numbered lists, code, pipe tables) feeds the three
writers so the files look alike.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from typing import Any
from xml.sax.saxutils import escape

import openpyxl
from docx import Document
from openpyxl.styles import Font
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

CONTENT_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "markdown": "text/markdown; charset=utf-8",
    "json": "application/json",
}
EXTENSIONS = {"docx": "docx", "xlsx": "xlsx", "pdf": "pdf", "markdown": "md", "json": "json"}
MAX_INPUT_TEXT = 200_000  # characters a file input can contribute to the graph variables
SHEET_NAME_BAD = re.compile(r"[\\/*?:\[\]]")


@dataclass
class Block:
    kind: str  # heading | paragraph | bullets | numbers | code | table
    text: str = ""
    level: int = 1
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def text_of(value: Any) -> str:
    """The final output as Markdown: text as is, anything else as a fenced JSON block."""
    if isinstance(value, str):
        return value
    if value is None:
        return ""
    return "```json\n" + json.dumps(value, indent=2, ensure_ascii=False, default=str) + "\n```"


def parse_blocks(markdown: str) -> list[Block]:
    lines = markdown.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    para: list[str] = []
    i = 0

    def flush() -> None:
        if para:
            blocks.append(Block("paragraph", text=" ".join(p.strip() for p in para)))
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(Block("code", text="\n".join(code)))
            i += 1
            continue
        if (m := re.match(r"^(#{1,6})\s+(.*)$", stripped)) is not None:
            flush()
            blocks.append(Block("heading", text=m.group(2).strip(), level=len(m.group(1))))
        elif "|" in stripped and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1]):
            flush()
            rows = [_cells(stripped)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_cells(lines[i]))
                i += 1
            blocks.append(Block("table", rows=rows))
            continue
        elif re.match(r"^[-*+]\s+", stripped) or re.match(r"^\d+[.)]\s+", stripped):
            flush()
            numbered = bool(re.match(r"^\d+[.)]\s+", stripped))
            items: list[str] = []
            while i < len(lines):
                s = lines[i].strip()
                m2 = re.match(r"^(?:[-*+]|\d+[.)])\s+(.*)$", s)
                if not m2 or bool(re.match(r"^\d+[.)]\s+", s)) != numbered:
                    break
                items.append(m2.group(1))
                i += 1
            blocks.append(Block("numbers" if numbered else "bullets", items=items))
            continue
        elif not stripped:
            flush()
        else:
            para.append(stripped)
        i += 1
    flush()
    return blocks


INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|`[^`]+`)")


def _runs(text: str) -> list[tuple[str, str]]:
    """Inline Markdown as (style, text): style is '', 'b', 'i' or 'c'."""
    out: list[tuple[str, str]] = []
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith(("**", "__")) and part.endswith(("**", "__")) and len(part) > 4:
            out.append(("b", part[2:-2]))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            out.append(("c", part[1:-1]))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append(("i", part[1:-1]))
        else:
            out.append(("", part))
    return out


def _plain(text: str) -> str:
    return "".join(t for _, t in _runs(text))


# ---------------------------------------------------------------------------- DOCX
def render_docx(title: str, value: Any) -> bytes:
    doc = Document()
    doc.add_heading(title, level=0)
    for block in parse_blocks(text_of(value)):
        if block.kind == "heading":
            doc.add_heading(_plain(block.text), level=min(block.level, 4))
        elif block.kind == "paragraph":
            _docx_runs(doc.add_paragraph(), block.text)
        elif block.kind in ("bullets", "numbers"):
            style = "List Bullet" if block.kind == "bullets" else "List Number"
            for item in block.items:
                _docx_runs(doc.add_paragraph(style=style), item)
        elif block.kind == "code":
            run = doc.add_paragraph().add_run(block.text)
            run.font.name = "Courier New"
        elif block.kind == "table":
            width = max(len(r) for r in block.rows)
            table = doc.add_table(rows=len(block.rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(block.rows):
                for c in range(width):
                    cell = table.cell(r, c)
                    cell.text = _plain(row[c]) if c < len(row) else ""
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_runs(paragraph: Any, text: str) -> None:
    for style, part in _runs(text):
        run = paragraph.add_run(part)
        run.bold = style == "b"
        run.italic = style == "i"
        if style == "c":
            run.font.name = "Courier New"


# ---------------------------------------------------------------------------- XLSX
def _sheet_name(name: str, used: set[str]) -> str:
    base = SHEET_NAME_BAD.sub(" ", name).strip()[:28] or "Sheet"
    candidate, n = base, 2
    while candidate.lower() in used:
        candidate = f"{base[:25]} {n}"
        n += 1
    used.add(candidate.lower())
    return candidate


def _table_rows(value: Any) -> list[list[Any]] | None:
    """A JSON array of objects (or of arrays) as sheet rows, header first."""
    if isinstance(value, list) and value and all(isinstance(v, dict) for v in value):
        keys: list[str] = []
        for item in value:
            keys += [k for k in item if k not in keys]
        return [keys] + [[_cell(item.get(k)) for k in keys] for item in value]
    if isinstance(value, list) and value and all(isinstance(v, list) for v in value):
        return [[_cell(c) for c in row] for row in value]
    return None


def _cell(value: Any) -> Any:
    if value is None or isinstance(value, bool | int | float | str):
        return value
    return json.dumps(value, ensure_ascii=False, default=str)


def render_xlsx(title: str, value: Any) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used: set[str] = set()

    def add(name: str, rows: list[list[Any]], header: bool = True) -> None:
        ws = wb.create_sheet(_sheet_name(name, used))
        for row in rows:
            ws.append(row)
        if header and rows:
            for cell in ws[1]:
                cell.font = Font(bold=True)
        for col in ws.columns:
            longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(longest + 2, 10), 60)

    parsed = value
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except ValueError:
            parsed = value
    rows = _table_rows(parsed)
    if rows is not None:
        add("Data", rows)
    elif isinstance(parsed, dict):
        add("Data", [["Field", "Value"]] + [[str(k), _cell(v)] for k, v in parsed.items()])
    else:
        text_blocks = parse_blocks(text_of(parsed))
        tables = [b for b in text_blocks if b.kind == "table"]
        for n, table in enumerate(tables, 1):
            add(f"Table {n}", [[_plain(c) for c in r] for r in table.rows])
        body: list[list[Any]] = [[title], []]
        for b in text_blocks:
            if b.kind in ("heading", "paragraph"):
                body.append([_plain(b.text)])
            elif b.kind == "code":
                body += [[line] for line in b.text.split("\n")]
            elif b.kind in ("bullets", "numbers"):
                body += [["• " + _plain(i)] for i in b.items]
        if not tables or len(body) > 2:
            add("Output", body, header=False)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------- PDF
def render_pdf(title: str, value: Any) -> bytes:
    styles = getSampleStyleSheet()
    story: list[Any] = [Paragraph(escape(title), styles["Title"]), Spacer(1, 0.4 * cm)]

    def inline(text: str) -> str:
        out = ""
        for style, part in _runs(text):
            part = escape(part)
            out += {
                "b": f"<b>{part}</b>",
                "i": f"<i>{part}</i>",
                "c": f"<font face='Courier'>{part}</font>",
            }.get(style, part)
        return out

    for block in parse_blocks(text_of(value)):
        if block.kind == "heading":
            story.append(Paragraph(inline(block.text), styles[f"Heading{min(block.level, 4)}"]))
        elif block.kind == "paragraph":
            story += [Paragraph(inline(block.text), styles["BodyText"]), Spacer(1, 0.15 * cm)]
        elif block.kind in ("bullets", "numbers"):
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(inline(i), styles["BodyText"])) for i in block.items],
                    bulletType="bullet" if block.kind == "bullets" else "1",
                )
            )
        elif block.kind == "code":
            story += [Preformatted(block.text, styles["Code"]), Spacer(1, 0.15 * cm)]
        elif block.kind == "table":
            width = max(len(r) for r in block.rows)
            data = [
                [
                    Paragraph(inline(r[c]) if c < len(r) else "", styles["BodyText"])
                    for c in range(width)
                ]
                for r in block.rows
            ]
            table = Table(data, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story += [table, Spacer(1, 0.2 * cm)]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=title, author="Forge").build(story)
    return buf.getvalue()


def render(kind: str, title: str, value: Any) -> bytes:
    if kind == "docx":
        return render_docx(title, value)
    if kind == "xlsx":
        return render_xlsx(title, value)
    if kind == "pdf":
        return render_pdf(title, value)
    if kind == "json":
        return json.dumps(value, indent=2, ensure_ascii=False, default=str).encode()
    return text_of(value).encode()


# ---------------------------------------------------------------------------- input files
def extract_text(name: str, data: bytes) -> str:
    """Text of an uploaded file: plain text formats as is, PDF, DOCX and XLSX read. Binary
    formats Forge cannot read raise ValueError."""
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)[:MAX_INPUT_TEXT]
    if ext == "docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            parts += [" | ".join(c.text for c in row.cells) for row in table.rows]
        return "\n".join(parts)[:MAX_INPUT_TEXT]
    if ext == "xlsx":
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"# {ws.title}")
            for row in ws.iter_rows(values_only=True):
                parts.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(parts)[:MAX_INPUT_TEXT]
    try:
        return data.decode("utf-8")[:MAX_INPUT_TEXT]
    except UnicodeDecodeError as exc:
        raise ValueError(f"'{name}' is not a text, PDF, DOCX or XLSX file") from exc
